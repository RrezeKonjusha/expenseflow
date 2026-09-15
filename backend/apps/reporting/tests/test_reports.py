import json
from datetime import timedelta
from decimal import Decimal as D
from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.expenses.models import ExpenseStatus
from apps.reporting.documents import ReportSnapshot
from tests.factories import MealFactory, TravelFactory

pytestmark = pytest.mark.django_db
R = "/api/v1/reports/"


@pytest.fixture
def data(world):
    now = timezone.now()
    MealFactory(
        employee=world.arta,
        project=world.alpha,
        amount=D("40"),
        status=ExpenseStatus.APPROVED,
        decided_by=world.besa,
        decided_at=now,
    )
    TravelFactory(
        employee=world.arta,
        project=world.alpha,
        amount=D("30"),
        status=ExpenseStatus.REIMBURSED,
        decided_by=world.besa,
        decided_at=now,
    )
    MealFactory(
        employee=world.blerim,
        project=world.alpha,
        amount=D("20"),
        status=ExpenseStatus.APPROVED,
        decided_by=world.driton,
        decided_at=now,
    )
    MealFactory(employee=world.blerim, project=world.alpha, amount=D("15"))  # draft
    return world


def criteria(**over):
    today = timezone.localdate()
    c = {
        "date_from": str(today - timedelta(days=30)),
        "date_to": str(today),
        "group_by": "department",
        "statuses": ["APPROVED", "REIMBURSED"],
    }
    c.update(over)
    return c


def test_dashboard_scoped_and_cached(data):
    admin = data.client(data.admin)
    d = admin.get(f"{R}dashboard/").data
    assert d["scope"] == "all" and d["counts"]["APPROVED"] == 2 and len(d["monthly"]) == 6
    assert data.client(data.besa).get(f"{R}dashboard/").data["scope"] == f"department:{data.eng.pk}"
    MealFactory(employee=data.arta, project=data.alpha, status=ExpenseStatus.SUBMITTED)
    assert admin.get(f"{R}dashboard/").data["counts"]["SUBMITTED"] == 0  # served from cache
    from apps.expenses.services import bump_dashboard_cache

    bump_dashboard_cache()
    assert admin.get(f"{R}dashboard/").data["counts"]["SUBMITTED"] == 1


@pytest.mark.parametrize("group_by", ["department", "project", "employee", "type", "status", "month"])
def test_run_report_all_groupings(data, group_by):
    r = data.client(data.admin).post(R + "run/", criteria(group_by=group_by), format="json")
    assert r.status_code == 200, r.data
    assert D(r.data["totals"]["total"]) == D("90.00") and r.data["totals"]["count"] == 3


def test_report_by_department_values(data):
    rows = data.client(data.admin).post(R + "run/", criteria(), format="json").data["rows"]
    assert {r["label"]: D(r["total"]) for r in rows} == {"Engineering": D("70.00"), "Sales": D("20.00")}


def test_report_is_scoped_for_manager(data):
    r = data.client(data.besa).post(R + "run/", criteria(), format="json")
    assert [row["label"] for row in r.data["rows"]] == ["Engineering"]


def test_report_validation(data):
    c = data.client(data.admin)
    assert c.post(R + "run/", criteria(group_by="password"), format="json").status_code == 400
    bad = criteria(date_from=str(timezone.localdate()), date_to=str(timezone.localdate() - timedelta(days=1)))
    assert c.post(R + "run/", bad, format="json").status_code == 400


def test_save_list_export_delete_snapshot(data):
    c = data.client(data.admin)
    r = c.post(R, {"name": "Q3 by department", "criteria": criteria()}, format="json")
    assert r.status_code == 201
    sid = r.data["id"]
    stored = ReportSnapshot.objects.get(pk=sid)
    assert stored.owner.email == "admin@test.dev" and len(stored.rows) == 2  # embedded documents
    assert [s["id"] for s in c.get(R).data] == [sid]
    assert c.get(f"{R}{sid}/").data["name"] == "Q3 by department"

    xlsx = c.get(f"{R}{sid}/export/", {"format": "xlsx"})
    assert xlsx.status_code == 200 and "spreadsheetml" in xlsx["Content-Type"]
    ws = load_workbook(BytesIO(xlsx.content)).active
    assert ws["A4"].value == "Group" and ws.cell(ws.max_row, 1).value == "TOTAL"

    csv_resp = c.get(f"{R}{sid}/export/", {"format": "csv"})
    assert b"Engineering" in csv_resp.content
    js = json.loads(c.get(f"{R}{sid}/export/", {"format": "json"}).content)
    assert js["totals"]["count"] == 3
    assert c.get(f"{R}{sid}/export/", {"format": "pdf"}).status_code == 400

    assert data.client(data.arta).get(f"{R}{sid}/").status_code == 404  # not the owner
    assert c.delete(f"{R}{sid}/").status_code == 204
    assert c.get(f"{R}{sid}/").status_code == 404
    assert c.get(f"{R}not-an-id/").status_code == 404


def test_csv_formula_injection_is_neutralised():
    from apps.reporting.services import _csv_safe

    assert _csv_safe("=SUM(A1)") == "'=SUM(A1)" and _csv_safe("Sales") == "Sales"
