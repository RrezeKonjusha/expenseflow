"""Statistics & reporting: dashboard KPIs, dynamic reports (extra feature 11), exports (feature 10)."""

import csv
import io
import json
from decimal import Decimal

from dateutil.relativedelta import relativedelta
from django.conf import settings
from django.core.cache import cache
from django.db.models import Count, F, Sum
from django.db.models.functions import Coalesce, TruncMonth
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font

from apps.core import audit
from apps.expenses.models import ExpenseStatus
from apps.expenses.selectors import expenses_visible_to
from apps.expenses.services import DASHBOARD_VERSION_KEY

from .documents import Criteria, Owner, ReportSnapshot, Row, Totals

SPENT = [ExpenseStatus.APPROVED, ExpenseStatus.REIMBURSED]
ZERO = Decimal("0.00")

# group_by whitelist: (ORM expression for the key, ORM field for the label)
GROUPS = {
    "department": ("employee__department_id", "employee__department__name"),
    "project": ("project_id", "project__code"),
    "employee": ("employee_id", "employee__email"),
    "type": ("type", "type"),
    "status": ("status", "status"),
    "month": ("month", "month"),
}


def scope_of(user) -> str:
    if user.is_admin:
        return "all"
    if user.is_manager and user.department_id:
        return f"department:{user.department_id}"
    return f"user:{user.pk}"


# ---- dashboard (cached) --------------------------------------------------------------


def dashboard(user) -> dict:
    version = cache.get_or_set(DASHBOARD_VERSION_KEY, 1, None)
    key = f"dashboard:v{version}:{scope_of(user)}"
    data = cache.get(key)
    if data is None:
        data = _compute_dashboard(user)
        cache.set(key, data, settings.DASHBOARD_CACHE_SECONDS)
    return data


def _compute_dashboard(user) -> dict:
    qs = expenses_visible_to(user, polymorphic=False)
    today = timezone.localdate()
    month_start = today.replace(day=1)
    six_months_ago = month_start - relativedelta(months=5)

    by_status = {r["status"]: r["n"] for r in qs.values("status").annotate(n=Count("id"))}
    spent_month = qs.filter(status__in=SPENT, expense_date__gte=month_start).aggregate(s=Coalesce(Sum("amount"), ZERO))[
        "s"
    ]
    pending_amount = qs.filter(status=ExpenseStatus.SUBMITTED).aggregate(s=Coalesce(Sum("amount"), ZERO))["s"]
    monthly = (
        qs.filter(status__in=SPENT, expense_date__gte=six_months_ago)
        .annotate(month=TruncMonth("expense_date"))
        .values("month")
        .annotate(total=Sum("amount"), count=Count("id"))
        .order_by("month")
    )
    by_month = {r["month"].strftime("%Y-%m"): r for r in monthly}
    series = []
    for i in range(6):
        m = (six_months_ago + relativedelta(months=i)).strftime("%Y-%m")
        r = by_month.get(m)
        series.append({"month": m, "total": str(r["total"] if r else ZERO), "count": r["count"] if r else 0})
    by_type = [
        {"type": r["type"], "total": str(r["total"])}
        for r in qs.filter(status__in=SPENT).values("type").annotate(total=Sum("amount")).order_by("type")
    ]
    return {
        "scope": scope_of(user),
        "counts": {s: by_status.get(s, 0) for s in ExpenseStatus.values},
        "spent_this_month": str(spent_month),
        "pending_amount": str(pending_amount),
        "monthly": series,
        "by_type": by_type,
        "generated_at": timezone.now().isoformat(),
    }


# ---- dynamic report ---------------------------------------------------------------------


def run_report(user, c: dict) -> dict:
    qs = expenses_visible_to(user, polymorphic=False).filter(expense_date__range=(c["date_from"], c["date_to"]))
    if c.get("statuses"):
        qs = qs.filter(status__in=c["statuses"])
    if c.get("types"):
        qs = qs.filter(type__in=c["types"])
    if c.get("department_ids"):
        qs = qs.filter(employee__department_id__in=c["department_ids"])
    if c.get("project_ids"):
        qs = qs.filter(project_id__in=c["project_ids"])

    key_expr, label_expr = GROUPS[c["group_by"]]
    if c["group_by"] == "month":
        qs = qs.annotate(month=TruncMonth("expense_date"))
        grouped = qs.values("month").annotate(count=Count("id"), total=Sum("amount")).order_by("month")
        rows = [
            {
                "key": r["month"].strftime("%Y-%m"),
                "label": r["month"].strftime("%b %Y"),
                "count": r["count"],
                "total": r["total"],
            }
            for r in grouped
        ]
    else:
        grouped = (
            qs.annotate(k=F(key_expr), lbl=F(label_expr))
            .values("k", "lbl")
            .annotate(count=Count("id"), total=Sum("amount"))
            .order_by("-total")
        )
        rows = [
            {
                "key": str(r["k"]) if r["k"] is not None else "none",
                "label": r["lbl"] or "(none)",
                "count": r["count"],
                "total": r["total"],
            }
            for r in grouped
        ]
    totals = {"count": sum(r["count"] for r in rows), "total": sum((r["total"] for r in rows), ZERO)}
    return {"criteria": c, "rows": rows, "totals": totals}


def save_snapshot(user, *, name: str, criteria: dict, request=None) -> ReportSnapshot:
    result = run_report(user, criteria)
    snap = ReportSnapshot(
        name=name,
        owner=Owner(user_id=user.pk, email=user.email),
        scope=scope_of(user),
        criteria=Criteria(**{k: v for k, v in criteria.items() if v is not None}),
        rows=[Row(**r) for r in result["rows"]],
        totals=Totals(**result["totals"]),
    )
    snap.save(write_concern={"w": "majority"} if not settings.MONGO_URL.startswith("mongomock") else None)
    audit.record("REPORT_SAVED", actor=user, target=("ReportSnapshot", snap.pk), request=request)
    return snap


def snapshots_visible_to(user):
    qs = ReportSnapshot.objects
    return qs if user.is_admin else qs.filter(owner__user_id=user.pk)


# ---- export ---------------------------------------------------------------------------------

HEADER = ["Group", "Key", "Count", "Total (EUR)"]


def export_snapshot(snap: ReportSnapshot, fmt: str) -> tuple[bytes, str, str]:
    base = f"report-{snap.pk}"
    rows = [[r.label, r.key, r.count, Decimal(str(r.total))] for r in snap.rows]
    if fmt == "json":
        payload = {
            "name": snap.name,
            "generated_at": snap.generated_at.isoformat(),
            "criteria": json.loads(snap.criteria.to_json()),
            "rows": [{"label": a, "key": b, "count": c, "total": str(d)} for a, b, c, d in rows],
            "totals": {"count": snap.totals.count, "total": str(snap.totals.total)},
        }
        return json.dumps(payload, indent=2, default=str).encode(), "application/json", f"{base}.json"
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(HEADER)
        for r in rows:
            w.writerow([_csv_safe(r[0]), r[1], r[2], f"{r[3]:.2f}"])
        w.writerow(["TOTAL", "", snap.totals.count, f"{Decimal(str(snap.totals.total)):.2f}"])
        return buf.getvalue().encode("utf-8-sig"), "text/csv", f"{base}.csv"
    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Report"
        ws.append([snap.name])
        ws["A1"].font = Font(bold=True, size=14)
        c = snap.criteria
        ws.append([f"{c.date_from} to {c.date_to}, grouped by {c.group_by}"])
        ws.append([])
        ws.append(HEADER)
        for cell in ws[4]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        ws.append(["TOTAL", "", snap.totals.count, Decimal(str(snap.totals.total))])
        ws[ws.max_row][0].font = Font(bold=True)
        for col, width in zip("ABCD", (30, 14, 10, 16), strict=True):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=5, min_col=4, max_col=4):
            for cell in row:
                cell.number_format = "#,##0.00"
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", f"{base}.xlsx"
    raise ValueError(f"Unsupported format {fmt}")


def _csv_safe(value: str) -> str:
    """Prevent CSV formula injection when the file is opened in Excel."""
    return "'" + value if value and value[0] in "=+-@" else value
